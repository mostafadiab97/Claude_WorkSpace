"""Generate a realistic, fictional requirements spreadsheet for ReqMan testing."""
import random
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

random.seed(42)
out = sys.argv[1]

owners = ["A. Hassan", "M. Farouk", "S. Adel", "N. Kamal", "Y. Saber", "R. Nabil"]
statuses = ["Draft", "Review", "Approved", "Obsolete"]
status_w = [35, 20, 40, 5]
prios = ["High", "Medium", "Low"]

stk = [
    ("Comfortable temperature", "The user wants the room to stay at a chosen temperature without manual adjustment."),
    ("Remote control", "The user wants to view and change the temperature from a mobile phone."),
    ("Energy saving", "The user wants lower energy bills without losing comfort."),
    ("Easy installation", "A non-specialist installer shall be able to fit the device in under 30 minutes."),
    ("Safe operation", "The device shall never create a fire or electrical hazard."),
    ("Schedules", "The user wants different temperatures at different times of day."),
    ("Status visibility", "The user wants to see current and target temperature at a glance."),
    ("Privacy", "The user's usage data shall not be shared without consent."),
]
sys_tpl = [
    ("Temperature measurement accuracy", "The thermostat shall measure room temperature with an accuracy of ±{a} °C.", ""),
    ("Relay switching latency", "The thermostat shall switch the heating relay within {t} ms of a control decision.", ""),
    ("Steady-state regulation", "The thermostat shall maintain the set temperature within ±{a} °C in steady state.", "ux"),
    ("Wi-Fi connectivity", "The thermostat shall connect to the home Wi-Fi network using WPA2 or WPA3.", "connectivity"),
    ("Remote set-point response", "The thermostat shall accept temperature changes from the mobile app within {s} s.", "connectivity"),
    ("Daily schedule capacity", "The thermostat shall support at least {n} schedule entries per day.", "ux"),
    ("Low-power idle mode", "The thermostat shall enter low-power mode after {m} minutes without input.", "power"),
    ("Display readability", "The display shall show current and target temperature with a minimum digit height of {h} mm.", "ux"),
    ("Over-temperature cut-off", "The thermostat shall cut power to the relay if internal temperature exceeds {c} °C.", "safety"),
    ("Schedule persistence", "The thermostat shall store schedules in non-volatile memory.", ""),
    ("Cloud data encryption", "The thermostat shall encrypt all data sent to the cloud service.", "security"),
    ("Firmware version reporting", "The thermostat shall report its firmware version to the mobile app.", "connectivity"),
]
sw_tpl = [
    ("Regulation loop period", "The control software shall run the regulation loop every {t} ms.", ""),
    ("App connection status", "The app software shall display connection status within {s} s of launch.", "ux"),
    ("Signed firmware updates", "The firmware shall validate update packages with a digital signature before installing.", "security"),
    ("Relay event logging", "The firmware shall log relay switching events with a timestamp.", ""),
    ("Schedule execution timing", "The scheduler shall apply the next schedule entry within {s} s of its start time.", ""),
    ("Power-loss recovery", "The firmware shall restore the last set temperature after a power loss.", "power"),
    ("Set-point range check", "The software shall reject set temperatures outside 5–30 °C.", "safety"),
    ("Analytics consent", "The app shall ask for consent before enabling usage analytics.", "security"),
]
hw_tpl = [
    ("Sensor resolution", "The temperature sensor shall have a resolution of 0.1 °C or better.", ""),
    ("Relay current rating", "The relay shall be rated for at least {r} A at 230 V AC.", "safety"),
    ("Enclosure protection", "The enclosure shall meet IP{ip} protection.", "safety"),
    ("Mains supply range", "The device shall operate from 100–240 V AC, 50/60 Hz.", "power"),
]

rows = []
def add(i, typ, title, desc, parent="", verified="", tags=None):
    rows.append({
        "ID": i, "Type": typ, "Title": title, "Description": desc,
        "Status": random.choices(statuses, status_w)[0],
        "Priority": random.choice(prios), "Owner": random.choice(owners),
        "Parent": parent, "Verified By": verified,
        "Tags": tags if tags is not None else "",
    })

stk_ids = []
for n, (t, d) in enumerate(stk, 1):
    i = f"STK-{n:03d}"; stk_ids.append(i); add(i, "Stakeholder Need", t, d)

def fill(s):
    return s.format(a=random.choice([0.2, 0.3, 0.5]), t=random.choice([100, 200, 500]),
                    s=random.choice([2, 3, 5]), n=random.choice([4, 6, 8]),
                    m=random.choice([1, 2, 5]), h=random.choice([8, 10, 12]),
                    c=random.choice([70, 80, 85]), r=random.choice([5, 10, 16]),
                    ip=random.choice([20, 30, 40]))

def short(desc):
    w = desc.replace("The ", "").split()
    return " ".join(w[:6]).rstrip(".,")

sys_ids = []
for n, (title, tpl, tag) in enumerate(sys_tpl, 1):
    i = f"SYS-{n:03d}"; sys_ids.append(i)
    add(i, "System Requirement", title, fill(tpl), parent=random.choice(stk_ids), tags=tag)
sw_ids = []
for n, (title, tpl, tag) in enumerate(sw_tpl, 1):
    i = f"SW-{n:03d}"; sw_ids.append(i)
    add(i, "Software Requirement", title, fill(tpl), parent=random.choice(sys_ids), tags=tag)
hw_ids = []
for n, (title, tpl, tag) in enumerate(hw_tpl, 1):
    i = f"HW-{n:03d}"; hw_ids.append(i)
    add(i, "Hardware Requirement", title, fill(tpl), parent=random.choice(sys_ids), tags=tag)

# Test cases verifying system/software/hardware requirements
targets = sys_ids + sw_ids + hw_ids
tc_ids = []
for n in range(1, 21):
    i = f"TC-{n:03d}"; tc_ids.append(i)
    tgt = random.choice(targets)
    add(i, "Test Case", f"Verify {tgt}", f"Set up the test bench, exercise the behaviour described in {tgt} and record the result against its acceptance criteria.")
    # record the verification on the target row
    for r in rows:
        if r["ID"] == tgt:
            r["Verified By"] = ", ".join(filter(None, [r["Verified By"], i]))

risks = [
    ("Overheating of relay", "Relay contacts may overheat under sustained high load."),
    ("Wi-Fi credential leak", "Credentials could be exposed if stored unencrypted."),
    ("Sensor drift", "Sensor accuracy may degrade over time, causing discomfort."),
    ("Failed firmware update", "A power cut during update could leave the device unusable."),
    ("Mis-wiring at install", "Incorrect wiring may damage the heating system."),
]
for n, (t, d) in enumerate(risks, 1):
    add(f"RSK-{n:03d}", "Risk", t, d, parent=random.choice(stk_ids), tags="safety")

crs = [
    ("Add cooling support", "Extend control logic to support air-conditioning units."),
    ("Increase schedule entries", "Customers request up to 12 schedule entries per day."),
]
for n, (t, d) in enumerate(crs, 1):
    add(f"CR-{n:03d}", "Change Request", t, d)

wb = Workbook()
ws = wb.active
ws.title = "Requirements"
cols = ["ID", "Type", "Title", "Description", "Status", "Priority", "Owner", "Parent", "Verified By", "Tags"]
ws.append(cols)
for r in rows:
    ws.append([r[c] for c in cols])
widths = [10, 22, 34, 80, 11, 10, 12, 10, 18, 13]
for idx, w in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(1, idx).column_letter].width = w
for row in ws.iter_rows(min_row=2):
    row[3].alignment = Alignment(wrap_text=True, vertical="top")
for c in ws[1]:
    c.font = Font(bold=True)
tab = Table(displayName="Requirements", ref=f"A1:J{len(rows)+1}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tab)
ws.freeze_panes = "A2"

info = wb.create_sheet("About")
for line in [
    ["ReqMan sample data"],
    ["Fictional requirements for a smart thermostat, generated for testing import/export."],
    ["Not real product data."],
    [f"Rows: {len(rows)}"],
]:
    info.append(line)
info["A1"].font = Font(bold=True, size=14)
info.column_dimensions["A"].width = 90

wb.save(out)
print(len(rows), "rows written to", out)
